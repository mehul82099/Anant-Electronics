"""Verbatim dump of every non-empty region of every sheet. READ-ONLY."""
import json, os, datetime, hashlib
from collections import defaultdict
import openpyxl

SRC = r"C:\Users\Mehul\Downloads\MOBILE SHOP CATLOG FOR PHONES\data\MOP LIST NEW 28.10.25.xlsx"
OUT = r"C:\Users\Mehul\Downloads\MOBILE SHOP CATLOG FOR PHONES\analysis"

def cell_str(v):
    if v is None: return ""
    if isinstance(v, (datetime.datetime, datetime.date)): return v.isoformat()
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip()

def col_letter(i):
    s = ""
    while i:
        i, r = divmod(i - 1, 26); s = chr(65 + r) + s
    return s

wb = openpyxl.load_workbook(SRC, data_only=True)
out = {}
for name in wb.sheetnames:
    ws = wb[name]
    rows = []
    for r in range(1, (ws.max_row or 0) + 1):
        cells = {}
        for c in range(1, (ws.max_column or 0) + 1):
            s = cell_str(ws.cell(row=r, column=c).value)
            if s:
                cells[col_letter(c)] = s
        if cells:
            rows.append({"row": r, "cells": cells})
    out[name] = {
        "sheet": name,
        "state": ws.sheet_state,
        "max_row": ws.max_row, "max_col": ws.max_column,
        "populated_rows": len(rows),
        "first_row": rows[0]["row"] if rows else None,
        "last_row": rows[-1]["row"] if rows else None,
        "rows": rows,
    }
    print(f"{name:<26} populated_rows={len(rows):<5} range={out[name]['first_row']}..{out[name]['last_row']}")

p = os.path.join(OUT, "raw_dump.json")
with open(p, "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=1)
print("wrote", p, os.path.getsize(p), "bytes")
